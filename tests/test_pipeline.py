"""Unit tests for the pure parts of the pipeline. Run: python3 tests/test_pipeline.py"""
import importlib.machinery, importlib.util, json, tempfile, unittest
from pathlib import Path

l = importlib.machinery.SourceFileLoader("mom", str(Path(__file__).parent.parent / "pipeline" / "mom"))
m = importlib.util.module_from_spec(importlib.util.spec_from_loader("mom", l)); l.exec_module(m)

class Tidy(unittest.TestCase):
    def test_spacing_and_punctuation(self):
        self.assertEqual(m.tidy_text("  pada renewal  SSL ,terdapat issue  "), "Pada renewal SSL, terdapat issue.")
    def test_keeps_terminal_punctuation(self):
        self.assertEqual(m.tidy_text("Is it done?"), "Is it done?")
    def test_names_dedupe_and_case(self):
        self.assertEqual(m.tidy_names(["adelio  ferhan", "ADELIO FERHAN", "Ivan Saputra ", ""]), ["Adelio Ferhan", "Ivan Saputra"])
    def test_field_without_sentence_rules(self):
        self.assertEqual(m.tidy_text("operational - update ssl", sentence=False), "operational - update ssl")

class Echo(unittest.TestCase):
    def test_drops_mic_echo_of_meeting_audio(self):
        segs = [(10.0, "Others", "We will submit the change request today"), (10.8, "Me", "we will submit the change request today"), (30.0, "Me", "Thanks, noted")]
        out = m.drop_echo(segs)
        self.assertEqual([x[2] for x in out], ["We will submit the change request today", "Thanks, noted"])
    def test_keeps_all_when_single_track(self):
        segs = [(1.0, "Speaker", "hello"), (5.0, "Speaker", "hello")]
        self.assertEqual(m.drop_echo(segs), segs)

class Loops(unittest.TestCase):
    def test_collapses_repetition_keeps_rest(self):
        self.assertEqual(m.collapse_loops("Hi in the chat in the chat in the chat in the chat. Hello everyone."), "Hi in the chat in the chat. Hello everyone.")
    def test_plain_text_untouched(self):
        self.assertEqual(m.collapse_loops("no loop here at all"), "no loop here at all")

class WeekdayDates(unittest.TestCase):
    def test_corrects_wrong_friday_and_saturday(self):
        import datetime
        d0 = datetime.date(2026, 10, 20)   # Tuesday
        self.assertEqual(m.fix_weekday_dates("paling lambat hari Jumat, 21 Oktober 2026", d0), "paling lambat hari Jumat, 23 Oktober 2026")
        self.assertEqual(m.fix_weekday_dates("pada Sabtu (22 Oktober 2026)", d0), "pada Sabtu (24 Oktober 2026)")
    def test_trusts_date_when_weekday_was_not_said(self):
        import datetime
        d0 = datetime.date(2026, 10, 20)
        self.assertEqual(m.fix_weekday_dates("pada Senin, 21 Oktober 2026", d0, {"jumat"}), "pada Rabu, 21 Oktober 2026")
        self.assertEqual(m.fix_weekday_dates("hari Jumat, 21 Oktober 2026", d0, {"jumat"}), "hari Jumat, 23 Oktober 2026")
    def test_keeps_consistent_dates(self):
        import datetime
        self.assertEqual(m.fix_weekday_dates("Rabu, 21 Oktober 2026", datetime.date(2026, 10, 20)), "Rabu, 21 Oktober 2026")

class Annotate(unittest.TestCase):
    def test_spells_out_relative_dates(self):
        import datetime
        out = m.annotate_dates("I will prepare it tomorrow. Targetnya jangan lewat Jumat, ya. Hari Sabtu malam.", datetime.date(2026, 10, 20))
        self.assertIn("tomorrow [Rabu 21 Oktober 2026]", out)
        self.assertIn("Jumat [Jumat 23 Oktober 2026]", out)
        self.assertIn("Sabtu [Sabtu 24 Oktober 2026]", out)

class NextWeek(unittest.TestCase):
    def test_weekday_depan_is_next_week(self):
        import datetime
        out = m.annotate_dates("deadline-nya Rabu depan, we meet next Friday", datetime.date(2026, 10, 20))
        self.assertIn("Rabu depan [Rabu 28 Oktober 2026]", out)
        self.assertIn("next Friday [Jumat 30 Oktober 2026]", out)

class Unsupported(unittest.TestCase):
    def test_drops_invented_weekday_date_keeps_supported(self):
        import datetime
        d0 = datetime.date(2026, 10, 20)
        tr = "mulai Kamis [Kamis 22 Oktober 2026] ya"
        out = m.drop_unsupported_dates(["yaitu hingga Minggu 25 Oktober 2026", "mulai Kamis 22 Oktober 2026"], tr, d0)
        self.assertEqual(out, ["yaitu hingga Minggu", "mulai Kamis 22 Oktober 2026"])

class MingguMeansWeek(unittest.TestCase):
    def test_minggu_as_week_is_not_sunday(self):
        import datetime
        out = m.annotate_dates("sudah empat minggu, dua minggu terakhir, laporan mingguan, minggu depan", datetime.date(2026, 10, 20))
        self.assertNotIn("Minggu 25", out)
        self.assertIn("minggu depan [week of", out)
    def test_hari_minggu_is_sunday(self):
        import datetime
        self.assertIn("hari Minggu [Minggu 25 Oktober 2026]", m.annotate_dates("hari Minggu kita libur", datetime.date(2026, 10, 20)))

class Naming(unittest.TestCase):
    def test_export_stem(self):
        d = Path("/x/2026-10-20_1500_mock2")
        self.assertEqual(m.export_stem({"agenda": "Operational – Update SSL Ondeck"}, d), "MOM_Operational Update SSL Ondeck_20261020")
        self.assertEqual(m.export_stem({"agenda": ""}, d), "MOM_Meeting_20261020")
        self.assertEqual(m.export_stem({"agenda": "Review: CTR / Budget?"}, d), "MOM_Review CTR Budget_20261020")

class Recurring(unittest.TestCase):
    def test_recurring_weekday_gets_no_date(self):
        import datetime
        out = m.annotate_dates("laporan tiap Senin pagi, setiap Jumat, every Monday", datetime.date(2026, 10, 20))
        self.assertNotIn("[", out)

class Dates(unittest.TestCase):
    def test_ordinals(self):
        self.assertEqual(m.nice_date(2026, 10, 19), "October 19th, 2026")
        self.assertEqual(m.nice_date(2026, 10, 21), "October 21st, 2026")
        self.assertEqual(m.nice_date(2026, 10, 22), "October 22nd, 2026")
        self.assertEqual(m.nice_date(2026, 10, 11), "October 11th, 2026")

class Segments(unittest.TestCase):
    def test_prepare_collects_device_change_segments_in_order(self):
        with tempfile.TemporaryDirectory() as t:
            d = Path(t)
            for n in ("mic.caf", "mic_2.caf", "system.caf", "mic_10.caf"): (d / n).write_bytes(b"x" * 5000)
            (d / "mode.txt").write_text("Online\n")
            _, tracks = m.prepare(d)
            lab = dict(tracks)
            self.assertEqual([p.name for p in lab["Me"]], ["mic.caf", "mic_2.caf", "mic_10.caf"])
            self.assertIn("Others", lab)

class Templates(unittest.TestCase):
    """The generic template is built and filled inside a scratch home, never the user's real ~/.mom."""
    def setUp(self):
        import tempfile
        self.t = Path(tempfile.mkdtemp()); (self.t / ".mom").mkdir()
        self.old = (m.HOME, m.TEMPLATE, m.TPL_JSON)
        m.HOME = self.t; m.TEMPLATE = self.t / ".mom/template.docx"; m.TPL_JSON = self.t / ".mom/template.json"
    def tearDown(self):
        m.HOME, m.TEMPLATE, m.TPL_JSON = self.old
    DATA = {"agenda": "Roadmap - Q4", "date": "October 20th, 2026", "time": "10:00 AM", "telkomsel_names": ["Dewi", "Budi"], "telkomsel_unit": "Product",
            "third_party_names": ["John"], "third_party_company": "Acme", "topics": ["First point.", "Second point.", "Third point."]}
    def test_generic_template_has_no_logo_and_no_telkomsel(self):
        import docx
        (self.t / ".mom/logo.png").write_bytes(b"old")
        m.make_generic_template(None, "")
        self.assertFalse((self.t / ".mom/logo.png").exists())
        d = docx.Document(str(m.TEMPLATE))
        self.assertNotIn("telkomsel", " ".join(c.text for r in d.tables[0].rows for c in r.cells).lower())
        self.assertEqual(len(d.inline_shapes), 0)
    def test_generic_template_uses_logo_when_given(self):
        import docx
        logo = self.t / "logo.png"
        import subprocess; subprocess.run(["/usr/bin/sips", "-s", "format", "png", "/System/Library/CoreServices/DefaultDesktop.heic", "--resampleWidth", "80", "--out", str(logo)], capture_output=True)
        if not logo.exists(): self.skipTest("no sample image")
        m.make_generic_template(logo, "Acme")
        self.assertEqual(len(docx.Document(str(m.TEMPLATE)).inline_shapes), 1)
        self.assertTrue((self.t / ".mom/logo.png").exists())
    def test_fill_generic(self):
        import docx
        tp = m.make_generic_template(None, "Acme")
        doc = docx.Document(str(m.TEMPLATE)); m.fill_template(doc, tp, self.DATA)
        cells = [c.text for r in doc.tables[0].rows for c in r.cells]
        for want in ("Roadmap - Q4", "October 20th, 2026", "10:00 AM", "Dewi\nBudi", "Product", "John", "Acme"):
            self.assertIn(want, cells)
        items = [p.text for p in doc.paragraphs if p.style.name == "List Number"]
        self.assertEqual(items, self.DATA["topics"])
    def test_fill_without_list_examples_adds_numbered_lines(self):
        import docx
        doc = docx.Document(); doc.add_paragraph("Topics"); 
        tp = {"map": {}, "topic_style": "List Paragraph", "topic_heading": 0}
        m.fill_template(doc, tp, self.DATA)
        self.assertEqual([p.text for p in doc.paragraphs][1:], ["1. First point.", "2. Second point.", "3. Third point."])
    def test_profile_defaults_when_nothing_learned(self):
        tp = m.template_profile()
        self.assertTrue(tp["generic"]); self.assertEqual(tp["team"], "a team")
    def test_hardware_tiers(self):
        self.assertIn(m.device_profile()["tier"], ("standard", "light"))


class TranscriptCleanup(unittest.TestCase):
    def test_drops_whisper_credit_lines_and_repeats(self):
        segs = [(0, "Others", "Sous-titrage Société Radio-Canada"), (5, "Me", "Ya halo selamat pagi"), (9, "Me", "Thanks for watching!"),
                (20, "Others", "ini kalimat yang sama lagi"), (30, "Others", "ini kalimat yang sama lagi"), (40, "Others", "ini kalimat yang sama lagi"), (50, "Others", "Okay"), (60, "Others", "Okay"), (70, "Others", "Okay")]
        self.assertEqual([x[2] for x in m.drop_hallucinations(segs)], ["Ya halo selamat pagi", "Okay", "Okay", "Okay"])
    def test_lines_are_timed_by_speech_not_by_silence(self):
        # one block 0-30 s, but the speaker only talked 0-5 s and 22-27 s (others were talking in between)
        out = m.split_lines(0, 30, "Me", "Kalimat pertama di awal. Kalimat kedua jauh di akhir rapat itu.", maxlen=30, runs=[[0, 5], [22, 27]])
        self.assertEqual(len(out), 2)
        self.assertLess(out[0][0], 5); self.assertGreaterEqual(out[1][0], 22)
    def test_tail_of_previous_utterance_does_not_pull_the_line_back(self):
        out = m.split_lines(6.0, 18.0, "Me", "Baik, Dewi, tolong eskalasi ke vendor hari ini ya.", runs=[[0, 6.3], [14.2, 18.2]])
        self.assertAlmostEqual(out[0][0], 14.2, places=1)
    def test_segment_placed_in_silence_moves_to_next_speech(self):
        runs = [[0, 6], [16, 20]]
        self.assertEqual(m.align_to_speech(6.1, 9.0, runs, floor=6.0), (16, 20))      # whisper put it right after the first, but it was spoken at 16 s
        self.assertEqual(m.align_to_speech(1.0, 5.0, runs), (1.0, 5.0))                 # already over speech: untouched
        self.assertEqual(m.align_to_speech(3.0, 4.0, []), (3.0, 4.0))
    def test_long_block_becomes_timed_sentence_lines(self):
        out = m.split_lines(0, 30, "Me", "Satu kalimat pendek. " + "Dua kalimat yang jauh lebih panjang dari yang pertama, " * 3 + "selesai.")
        self.assertGreaterEqual(len(out), 2)
        self.assertEqual(out[0][0], 0)
        self.assertTrue(all(a[0] < b[0] for a, b in zip(out, out[1:])))


class Names(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.t = Path(tempfile.mkdtemp())
        self.old = (m.HOME, m.USER_TERMS, m.LEARNED, m.GLOSS)
        m.HOME = self.t; m.USER_TERMS = self.t / "terms.txt"; m.LEARNED = self.t / "corrections.txt"; m.GLOSS = self.t / "glossary.txt"
    def tearDown(self):
        m.HOME, m.USER_TERMS, m.LEARNED, m.GLOSS = self.old
    def test_lexicon_has_brands_places_and_user_terms(self):
        m.USER_TERMS.write_text("Kopi Anjing\nbakso = Bakso Pak Kumis\n")
        lex = m.lexicon()
        for want in ("Teazzi", "Indomaret", "Kopi Kenangan", "Sudirman", "Kopi Anjing"): self.assertIn(want, lex)
        self.assertNotIn("bakso = Bakso Pak Kumis", lex)
    def test_apply_corrections(self):
        self.assertEqual(m.apply_corrections("kasih TZ ya, tz kecil, ATZ", {"TZ": "Teazzi"}), "kasih Teazzi ya, tz kecil, ATZ")
        self.assertEqual(m.apply_corrections("ke indo maret dulu", {"Indo Maret": "Indomaret"}), "ke Indomaret dulu")
    def test_learns_recurring_terms_across_meetings(self):
        m.SEEN = self.t / "seen.json"
        for name, tr in (("2026-10-01_0900_a", "[00:00:01] Me: Besok kita deploy Kopikat ke production. Nanti Kopikat dicek ya"), ("2026-10-02_0900_b", "[00:00:01] Me: Hari ini review Kopikat lagi sama Budi")):
            d = self.t / name; d.mkdir(); (d / "transcript.md").write_text("# Transcript\n\n" + tr + "\n"); m.harvest_terms(d)
            if name.endswith("_a"): self.assertNotIn("Kopikat", m.glossary())     # once is not enough
        self.assertIn("Kopikat", m.glossary())
        self.assertNotIn("Besok", m.glossary())                               # sentence-initial word, also seen lowercase
    def test_learns_word_swaps_from_user_edits(self):
        d = self.t / "m"; d.mkdir()
        (d / "transcript.orig.md").write_text("# Transcript\n\n[00:00:01] Me: Kita kasih TZ ya sama pak Budi\n[00:00:09] Me: Okay lanjut\n")
        (d / "transcript.md").write_text("# Transcript\n\n[00:00:01] Me: Kita kasih Teazzi ya sama Pak Budi\n[00:00:09] Me: Okay lanjut\n")
        m.learn_from_edits(d)
        self.assertEqual(m.corrections().get("TZ"), "Teazzi")
        self.assertNotIn("pak", " ".join(m.corrections()).lower())       # case-only change is not a correction
        self.assertIn("Teazzi", m.glossary())
        self.assertEqual(m.apply_corrections("kasih TZ lagi"), "kasih Teazzi lagi")


class Context(unittest.TestCase):
    def test_calendar_and_screen_text_reach_the_prompt(self):
        import tempfile
        d = Path(tempfile.mkdtemp()); (d / "meeting.json").write_text(json.dumps({"title": "Weekly Sync Ads", "attendees": ["Budi", "Siti"]}))
        (d / "screens.json").write_text(json.dumps([{"t": 75, "text": "Q4 Budget  Teazzi campaign  IDR 450 juta"}, {"t": 90, "text": "  "}]))
        c = m.context_block(d)
        self.assertIn("Weekly Sync Ads", c); self.assertIn("Budi", c); self.assertIn("[00:01:15]", c); self.assertIn("IDR 450 juta", c)
        self.assertEqual(len(m.screen_events(d)), 1)
        self.assertIn("Teazzi", m.screen_terms(d))
    def test_user_flags_reach_the_prompt_with_what_was_said(self):
        import tempfile
        d = Path(tempfile.mkdtemp())
        (d / "transcript.md").write_text("# Transcript\n\n[00:00:10] Me: Selamat pagi.\n[00:01:00] Others: Dewi akan kirim kontrak ke vendor.\n[00:05:00] Me: Lanjut.\n")
        (d / "marks.json").write_text(json.dumps([{"t": 65, "kind": "action"}]))
        c = m.context_block(d)
        self.assertIn("ACTION", c); self.assertIn("[00:01:05]", c); self.assertIn("kirim kontrak", c); self.assertNotIn("Lanjut", c)
    def test_no_context_files_is_fine(self):
        import tempfile
        self.assertEqual(m.context_block(Path(tempfile.mkdtemp())), "")
    def test_diagnostics_hide_home_and_name(self):
        out = m.diagnostics()
        self.assertNotIn(str(Path.home()), out); self.assertIn("Minute diagnostics", out)


class FreshMac(unittest.TestCase):
    def test_doctor_runs_on_a_home_with_nothing_in_it(self):
        import subprocess, tempfile, os
        h = tempfile.mkdtemp()
        env = {**os.environ, "HOME": h}
        r = subprocess.run(["/usr/bin/python3", str(Path(__file__).parent.parent / "pipeline" / "mom"), "--doctor"], capture_output=True, text=True, env=env)
        self.assertEqual(r.returncode, 0, r.stderr[-300:])
        d = json.loads(r.stdout.split("DOCTOR ", 1)[1])
        self.assertFalse(next(i for i in d["items"] if i["id"] == "whisper")["ok"])      # no model in an empty home
        self.assertIn(d["tier"], ("standard", "light"))

if __name__ == "__main__":
    unittest.main()
